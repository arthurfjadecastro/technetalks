"""Gera outputs/AI-002/planejamento-evento.xlsx a partir de DEC-009.

Uso: python scripts/evento/build_planilha.py  (requer: pip install openpyxl)
Depois de gerada, a planilha pode ser editada à mão; rodar de novo sobrescreve as edições.
"""

from datetime import date, time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SAIDA = Path(__file__).resolve().parents[2] / "outputs" / "AI-002" / "planejamento-evento.xlsx"
EVENTO = date(2026, 10, 17)
RESUMO = (
    "Sábado, 17/10/2026 · chegada 10h30, início 11h, fim 17h · Gama/DF · "
    "11 a 20 pessoas no total · inscrições até 10/10 · previsão R$ 500"
)
STATUS = ["A fazer", "Em andamento", "Feito", "Não se aplica"]

CABECALHO = Font(bold=True, color="FFFFFF")
FUNDO_CABECALHO = PatternFill("solid", fgColor="1F4E78")
EDITAVEL = PatternFill("solid", fgColor="FFF2CC")
TITULO = Font(bold=True, size=14)
QUEBRA = Alignment(wrap_text=True, vertical="top")


def d(dia, mes):
    return date(2026, mes, dia)


# (frente, tarefa, responsável, prazo, observação)
CHECKLIST = [
    ("Espaço", "Desenhar a planta em 3 versões (seca, com móveis, com móveis e pessoas) com os 20 lugares e disponibilizar", "Henrique", d(20, 9), "Telas → sofá → 2 mesas com cadeiras → 2 mesas de plástico extras; quadro de vidro móvel; lanche e café embaixo da escada."),
    ("Espaço", "Definir a tela principal: projetor e telão, ou TV", "Arthur + Henrique", d(27, 9), "TV é o plano B. Sem projetor garantido até esta data, a TV vira a principal."),
    ("Espaço", "Separar as mesas de plástico extras, se a planta precisar", "Henrique", d(4, 10), ""),
    ("Espaço", "Testar notebook → cabo/adaptador → tela com material real", "Arthur + Henrique", d(4, 10), ""),
    ("Espaço", "Posicionar 2 ou 3 ventiladores cruzados na sala e testar", "Henrique", d(16, 10), ""),
    ("Palestras", "Enviar o questionário aos dois palestrantes", "A definir", d(16, 9), "docs/evento/QUESTIONARIO_PALESTRANTES.md"),
    ("Palestras", "Receber títulos e resumos das duas palestras", "A definir", d(21, 9), "Base dos flyers 2 e 3."),
    ("Palestras", "Receber slides e uma pergunta para o quiz de cada palestrante", "A definir", d(14, 10), ""),
    ("Palestras", "Ensaio com a tela escolhida", "Arthur + Henrique + palestrantes", d(16, 10), ""),
    ("Divulgação", "Definir como será a inscrição (link ou lista)", "A definir", d(20, 9), ""),
    ("Divulgação", "Fazer a arte dos flyers a partir dos textos", "A definir", d(20, 9), "docs/evento/DIVULGACAO.md. Flyer 1 primeiro."),
    ("Divulgação", "Adicionar as pessoas ao grupo do WhatsApp", "A definir", d(21, 9), ""),
    ("Divulgação", "Primeiro disparo: nome, foto e descrição do grupo + flyer 1", "A definir", d(21, 9), "Abre as inscrições."),
    ("Divulgação", "Flyer 2 — palestra do VIRUS", "A definir", d(26, 9), ""),
    ("Divulgação", "Flyer 3 — palestra do Coatio", "A definir", d(1, 10), ""),
    ("Divulgação", "Flyer 4 — vagas acabando", "A definir", d(6, 10), ""),
    ("Inscrições", "Avaliar pedidos de convidados de fora do grupo", "Arthur + Henrique", d(10, 10), "A pessoa vem com alguém do grupo e também se inscreve."),
    ("Inscrições", "Encerrar as inscrições e fechar o quantitativo confirmado", "Arthur + Henrique", d(10, 10), "Anotar o número na aba Cardápio."),
    ("Divulgação", "Enviar lembrete privado aos inscritos com o endereço", "A definir", d(15, 10), ""),
    ("Registro", "Convidar Guilherme Reis para fotos e vídeos, com o flyer", "Henrique", d(21, 9), "A confirmar até ele aceitar."),
    ("Compras", "Banner: arte e encomenda", "A definir", d(4, 10), "Previsão R$ 60."),
    ("Compras", "Brinde: comprar 2 canecas", "A definir", d(4, 10), "Previsão R$ 70."),
    ("Alimentação", "Confirmar quem prepara e serve o almoço no dia", "Henrique", d(4, 10), ""),
    ("Alimentação", "Definir as proporções do almoço para 11, 16 e 20 pessoas", "Henrique", d(4, 10), "Aba Cardápio."),
    ("Alimentação", "Levantar o pedido da padaria para 11, 16 e 20 pessoas", "Henrique", d(4, 10), "Pãezinhos, assados e café."),
    ("Alimentação", "Encomendar o lanche na padaria com o quantitativo confirmado", "Henrique", d(12, 10), "Só depois de 10/10."),
    ("Alimentação", "Comprar ingredientes do almoço, refrigerantes, sucos e café", "Henrique", d(16, 10), "Só depois de 10/10."),
    ("Dinâmicas", "Montar o quiz do QR code e o formulário de feedback", "A definir", d(14, 10), "docs/evento/SORTEIO_E_FEEDBACK.md"),
    ("Dinâmicas", "Preparar a dinâmica de integração da recepção", "Arthur + Henrique", d(14, 10), "Aproximar quem estiver isolado e formar grupos."),
    ("Dinâmicas", "Preparar a pergunta de abertura da Mesa 360", "Arthur + Henrique", d(14, 10), ""),
    ("Pós-evento", "Agradecer no grupo e compartilhar fotos e materiais autorizados", "A definir", d(20, 10), ""),
    ("Pós-evento", "Fechar o balanço com comprovantes e acertar a divisão", "Arthur + Henrique", d(24, 10), "Aba Orçamento."),
]

# (início, fim, bloco, responsável, operação)
ROTEIRO = [
    (time(10, 30), time(11, 0), "Chegada", "Arthur + Henrique", "Horário divulgado. Equipe já arrumou o espaço; água disponível."),
    (time(11, 0), time(12, 30), "Recepção e integração", "Arthur + Henrique", "Check-in, apresentação do dia e das dinâmicas; aproximar quem estiver isolado e formar grupos."),
    (time(12, 30), time(14, 0), "Almoço", "Henrique + preparo", "Das 13h30 às 13h45 a equipe prepara o equipamento do VIRUS."),
    (time(14, 0), time(14, 45), "Palestra 1 — VIRUS", "VIRUS", "IA no dia a dia. Proposta: 30 min de conteúdo + 15 de perguntas."),
    (time(14, 45), time(15, 0), "Intervalo", "Arthur + Henrique", "Troca de equipamento e preparação da palestra 2."),
    (time(15, 0), time(15, 45), "Palestra 2 — Coatio", "Coatio", "Design thinking, UI/UX e demonstração do CARANGA."),
    (time(15, 45), time(16, 15), "Lanche", "Henrique", "Embaixo da escada, com café."),
    (time(16, 15), time(17, 0), "Mesa 360, fechamento e sorteio", "Arthur + Henrique", "Sugestão: Mesa 360 até 16h45; quiz, sorteio das 2 canecas e feedback até 17h."),
]

# (categoria, previsão, quando gastar)
ORCAMENTO = [
    ("Banner", 60, "Agora"),
    ("Brinde (2 canecas)", 70, "Agora"),
    ("Almoço", 250, "Depois de 10/10"),
    ("Lanche", 200, "Depois de 10/10"),
    ("Bebidas", 70, "Depois de 10/10"),
    ("Café", 30, "Depois de 10/10"),
]

ALMOCO = ["Lasanha", "Estrogonofe", "Arroz", "Batata palha", "Salada", "Refrigerante", "Suco"]
LANCHE = ["Pãezinhos", "Assados", "Café"]


def cabecalho(ws, linha, colunas):
    for i, texto in enumerate(colunas, start=1):
        c = ws.cell(row=linha, column=i, value=texto)
        c.font, c.fill, c.alignment = CABECALHO, FUNDO_CABECALHO, QUEBRA


def larguras(ws, valores):
    for letra, largura in zip("ABCDEFGH", valores):
        ws.column_dimensions[letra].width = largura


def titulo(ws, texto, subtitulo):
    ws["A1"] = texto
    ws["A1"].font = TITULO
    ws["A2"] = subtitulo
    ws["A2"].font = Font(italic=True, color="555555")


def aba_checklist(wb):
    ws = wb.active
    ws.title = "Checklist"
    titulo(ws, "Technetalks — checklist", RESUMO)
    ws["A3"] = "Amarelo: preencher. Prazo vencido sem 'Feito' fica vermelho. Datas de divulgação são propostas até o aceite."
    cabecalho(ws, 4, ["Nº", "Frente", "Tarefa", "Responsável", "Prazo", "Status", "Observação"])
    for n, (frente, tarefa, resp, prazo, obs) in enumerate(CHECKLIST, start=1):
        linha = 4 + n
        valores = [n, frente, tarefa, resp, prazo, "A fazer", obs]
        for col, v in enumerate(valores, start=1):
            c = ws.cell(row=linha, column=col, value=v)
            c.alignment = QUEBRA
        ws.cell(row=linha, column=5).number_format = "dd/mm/yyyy"
        for col in (4, 6):
            ws.cell(row=linha, column=col).fill = EDITAVEL
    ultima = 4 + len(CHECKLIST)
    faixa = f"A5:G{ultima}"
    status = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=False)
    status.add(f"F5:F{ultima}")
    ws.add_data_validation(status)
    ws.conditional_formatting.add(faixa, FormulaRule(formula=['$F5="Feito"'], fill=PatternFill("solid", fgColor="D9EAD3")))
    ws.conditional_formatting.add(faixa, FormulaRule(formula=['$F5="Não se aplica"'], font=Font(color="999999")))
    ws.conditional_formatting.add(
        faixa,
        FormulaRule(formula=['AND($E5<TODAY(),$F5<>"Feito",$F5<>"Não se aplica")'], font=Font(color="C00000", bold=True)),
    )
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:G{ultima}"
    larguras(ws, [5, 13, 60, 22, 12, 15, 50])


def aba_roteiro(wb):
    ws = wb.create_sheet("Roteiro")
    titulo(ws, "Roteiro do dia — 17/10", "Equipe arruma o espaço antes das 10h30 e desmonta depois das 17h.")
    cabecalho(ws, 4, ["Início", "Fim", "Minutos", "Bloco", "Responsável", "Operação"])
    for i, (inicio, fim, bloco, resp, oper) in enumerate(ROTEIRO):
        linha = 5 + i
        ws.cell(row=linha, column=1, value=inicio).number_format = "hh:mm"
        ws.cell(row=linha, column=2, value=fim).number_format = "hh:mm"
        ws.cell(row=linha, column=3, value=f"=ROUND((B{linha}-A{linha})*1440,0)")
        for col, v in ((4, bloco), (5, resp), (6, oper)):
            ws.cell(row=linha, column=col, value=v).alignment = QUEBRA
    total = 5 + len(ROTEIRO)
    ws.cell(row=total, column=2, value="Total").font = Font(bold=True)
    ws.cell(row=total, column=3, value=f"=SUM(C5:C{total - 1})").font = Font(bold=True)
    ws.cell(row=total + 2, column=1, value="Se atrasar, encurte a Mesa 360; preserve o almoço e as palestras.")
    ws.freeze_panes = "A5"
    larguras(ws, [9, 9, 10, 30, 22, 70])


def aba_orcamento(wb):
    ws = wb.create_sheet("Orçamento")
    titulo(ws, "Orçamento e balanço (R$)", "R$ 500 é previsão, não teto. Gastem, guardem os comprovantes e fechem o balanço depois do evento.")
    cabecalho(ws, 4, ["Categoria", "Previsão", "Gasto real", "Quando gastar", "Pago por", "Comprovante / observação"])
    for i, (cat, prev, quando) in enumerate(ORCAMENTO):
        linha = 5 + i
        ws.cell(row=linha, column=1, value=cat)
        ws.cell(row=linha, column=2, value=prev)
        ws.cell(row=linha, column=4, value=quando)
        for col in (3, 5, 6):
            ws.cell(row=linha, column=col).fill = EDITAVEL
    fim = 4 + len(ORCAMENTO)
    pago = DataValidation(type="list", formula1='"Arthur,Henrique"', allow_blank=True)
    pago.add(f"E5:E{fim}")
    ws.add_data_validation(pago)
    t = fim + 1
    ws.cell(row=t, column=1, value="Soma").font = Font(bold=True)
    ws.cell(row=t, column=2, value=f"=SUM(B5:B{fim})").font = Font(bold=True)
    ws.cell(row=t, column=3, value=f"=SUM(C5:C{fim})").font = Font(bold=True)

    b = t + 2
    ws.cell(row=b, column=1, value="Balanço").font = TITULO
    linhas = [
        ("Previsão da edição", 500),
        ("Soma das categorias − previsão", f"=B{t}-B{b + 1}"),
        ("Gasto real total", f"=C{t}"),
        ("Pago por Arthur", f'=SUMIF(E5:E{fim},"Arthur",C5:C{fim})'),
        ("Pago por Henrique", f'=SUMIF(E5:E{fim},"Henrique",C5:C{fim})'),
        ("Parte de cada um (metade)", f"=C{t}/2"),
        ("Acerto: positivo = Henrique repassa a Arthur", f"=B{b + 4}-B{b + 6}"),
    ]
    for i, (rotulo, valor) in enumerate(linhas, start=1):
        ws.cell(row=b + i, column=1, value=rotulo)
        ws.cell(row=b + i, column=2, value=valor)
    for linha in range(5, b + len(linhas) + 1):
        for col in (2, 3):
            ws.cell(row=linha, column=col).number_format = '"R$" #,##0.00'
    larguras(ws, [42, 14, 14, 17, 12, 45])


def aba_cardapio(wb):
    ws = wb.create_sheet("Cardápio")
    titulo(ws, "Cardápio por cenário", "Henrique preenche as quantidades. Compra só depois de 10/10, com o quantitativo confirmado.")
    ws["A4"] = "Quantitativo confirmado em 10/10"
    ws["B4"].fill = EDITAVEL
    ws["A5"] = "Cenário para a compra"
    ws["B5"] = '=IF(B4="","",IF(B4<=11,11,IF(B4<=16,16,20)))'
    linha = 7
    for secao, itens in (("Almoço", ALMOCO), ("Lanche (padaria)", LANCHE)):
        ws.cell(row=linha, column=1, value=secao).font = Font(bold=True, size=12)
        cabecalho(ws, linha + 1, ["Item", "11 pessoas", "16 pessoas", "20 pessoas", "Unidade", "Observação"])
        for i, item in enumerate(itens):
            r = linha + 2 + i
            ws.cell(row=r, column=1, value=item)
            for col in range(2, 7):
                ws.cell(row=r, column=col).fill = EDITAVEL
            if item == "Lasanha":
                ws.cell(row=r, column=2, value=1)
                ws.cell(row=r, column=5, value="travessa")
        linha += len(itens) + 4
    larguras(ws, [34, 12, 12, 12, 12, 40])


def conferir():
    for a, b in zip(ROTEIRO, ROTEIRO[1:]):
        assert a[1] == b[0], f"buraco no roteiro entre {a[2]} e {b[2]}"
    assert ROTEIRO[0][0] == time(10, 30) and ROTEIRO[-1][1] == time(17, 0)
    lanche = next(r for r in ROTEIRO if r[2] == "Lanche")
    assert (lanche[0], lanche[1]) == (time(15, 45), time(16, 15))
    assert sum(v for _, v, _ in ORCAMENTO) == 680
    for _, tarefa, _, prazo, _ in CHECKLIST:
        assert prazo <= EVENTO or tarefa.startswith(("Agradecer", "Fechar o balanço")), tarefa
    comida = [p for f, t, _, p, _ in CHECKLIST if f == "Alimentação" and t.startswith(("Comprar", "Encomendar"))]
    assert len(comida) == 2
    assert all(p > d(10, 10) for p in comida), "comida só depois do quantitativo de 10/10"


def main():
    conferir()
    wb = Workbook()
    aba_checklist(wb)
    aba_roteiro(wb)
    aba_orcamento(wb)
    aba_cardapio(wb)
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SAIDA)
    print(f"{SAIDA} · {len(CHECKLIST)} tarefas · {SAIDA.stat().st_size // 1024} KB")


if __name__ == "__main__":
    main()
